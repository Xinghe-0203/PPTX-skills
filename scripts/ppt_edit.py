"""Deterministic in-place edits for existing PPTX files."""

from __future__ import annotations

import argparse
import json
from collections import Counter
from pathlib import Path
from typing import Any, Iterable

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

from ppt_project import create_backup
from pptx_helper import THEMES
from reference_ppt import (
    _body_candidates,
    _shape_font_data,
    _title_candidate,
    replace_image,
    replace_text,
)


def _slide_at(prs: Presentation, slide_index: int) -> Any:
    position = 0 if slide_index == 0 else slide_index - 1
    if not 0 <= position < len(prs.slides):
        raise IndexError(f"Slide index out of range: {slide_index}")
    return prs.slides[position]


def _save(prs: Presentation, path: str | Path) -> None:
    target = Path(path).resolve()
    create_backup(target)
    prs.save(str(target))


def _replace_cross_run(paragraph: Any, find: str, replacement: str) -> int:
    count = 0
    while find:
        runs = list(paragraph.runs)
        full = "".join(run.text for run in runs)
        offsets = []
        cursor = 0
        for run in runs:
            offsets.append((cursor, cursor + len(run.text)))
            cursor += len(run.text)
        match = None
        search_from = 0
        while True:
            start = full.find(find, search_from)
            if start < 0:
                break
            end = start + len(find)
            start_index = next(i for i, (_, right) in enumerate(offsets) if start < right)
            end_index = next(i for i, (_, right) in enumerate(offsets) if end <= right)
            if start_index != end_index:
                match = (start, end, start_index, end_index)
                break
            search_from = start + 1
        if match is None:
            break
        start, end, start_index, end_index = match
        start_left = offsets[start_index][0]
        end_left = offsets[end_index][0]
        prefix = runs[start_index].text[:start - start_left]
        suffix = runs[end_index].text[end - end_left:]
        runs[start_index].text = prefix + replacement
        for index in range(start_index + 1, end_index):
            runs[index].text = ""
        runs[end_index].text = suffix
        count += 1
    return count


def edit_text(pptx_path: str | Path, slide_index: int, find: str, replace: str) -> int:
    if not find:
        raise ValueError("find must not be empty")
    prs = Presentation(str(pptx_path))
    slide = _slide_at(prs, slide_index)
    replacements = 0
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if find in run.text:
                    occurrences = run.text.count(find)
                    run.text = run.text.replace(find, replace)
                    replacements += occurrences
            if find in "".join(run.text for run in paragraph.runs):
                replacements += _replace_cross_run(paragraph, find, replace)
    if replacements:
        _save(prs, pptx_path)
    return replacements


def _role_shapes(slide: Any, role: str) -> list[Any]:
    role = role.casefold()
    title = _title_candidate(slide)
    if role == "title":
        return [title] if title is not None else []
    bodies = _body_candidates(slide, title)
    if role in {"body", "content"}:
        return bodies[:1]
    if role == "subtitle":
        return bodies[:1]
    if role == "kicker":
        candidates = []
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False) or not shape.text.strip():
                continue
            _, sizes, _ = _shape_font_data(shape)
            if sizes and shape.top < Inches(1.9) and max(sizes, default=99) <= 16:
                candidates.append(shape)
        return candidates[:1]
    raise ValueError(f"Unsupported text role: {role}")


def edit_text_by_role(
    pptx_path: str | Path,
    slide_index: int,
    role: str,
    new_text: str,
) -> int:
    prs = Presentation(str(pptx_path))
    slide = _slide_at(prs, slide_index)
    shapes = _role_shapes(slide, role)
    for shape in shapes:
        replace_text(shape, text=new_text)
    if shapes:
        _save(prs, pptx_path)
    return len(shapes)


def edit_title(pptx_path: str | Path, slide_index: int, new_title: str) -> int:
    return edit_text_by_role(pptx_path, slide_index, "title", new_title)


def edit_kicker(pptx_path: str | Path, slide_index: int, new_kicker: str) -> int:
    return edit_text_by_role(pptx_path, slide_index, "kicker", new_kicker)


def swap_image(
    pptx_path: str | Path,
    slide_index: int,
    image_index: int,
    new_path: str | Path,
) -> bool:
    image = Path(new_path).resolve()
    if not image.exists():
        raise FileNotFoundError(image)
    prs = Presentation(str(pptx_path))
    slide = _slide_at(prs, slide_index)
    pictures = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    position = 0 if image_index == 0 else image_index - 1
    if not 0 <= position < len(pictures):
        raise IndexError(f"Image index out of range: {image_index}")
    replace_image(slide, pictures[position], image)
    _save(prs, pptx_path)
    return True


def _walk_shapes(shapes: Iterable[Any]) -> Iterable[Any]:
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk_shapes(shape.shapes)


def _color_hex(color_format: Any) -> str | None:
    try:
        rgb = color_format.rgb
    except (AttributeError, TypeError, ValueError):
        return None
    return f"#{str(rgb).upper()}" if rgb is not None else None


def _replace_color_format(color_format: Any, mapping: dict[str, str]) -> int:
    current = _color_hex(color_format)
    if current not in mapping:
        return 0
    color_format.rgb = RGBColor.from_string(mapping[current].lstrip("#"))
    return 1


def _recolor_prs(prs: Presentation, mapping: dict[str, str]) -> int:
    normalized = {old.upper(): new.upper() for old, new in mapping.items()}
    replaced = 0
    for slide in prs.slides:
        try:
            replaced += _replace_color_format(slide.background.fill.fore_color, normalized)
        except (AttributeError, TypeError, ValueError):
            pass
        for shape in _walk_shapes(slide.shapes):
            try:
                replaced += _replace_color_format(shape.fill.fore_color, normalized)
            except (AttributeError, TypeError, ValueError):
                pass
            try:
                replaced += _replace_color_format(shape.line.color, normalized)
            except (AttributeError, TypeError, ValueError):
                pass
            if getattr(shape, "has_text_frame", False):
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        replaced += _replace_color_format(run.font.color, normalized)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        try:
                            replaced += _replace_color_format(cell.fill.fore_color, normalized)
                        except (AttributeError, TypeError, ValueError):
                            pass
                        for paragraph in cell.text_frame.paragraphs:
                            for run in paragraph.runs:
                                replaced += _replace_color_format(run.font.color, normalized)
    return replaced


def recolor(pptx_path: str | Path, old_hex: str, new_hex: str) -> int:
    for value in (old_hex, new_hex):
        if not value.startswith("#") or len(value) != 7:
            raise ValueError(f"Expected #RRGGBB color: {value}")
    prs = Presentation(str(pptx_path))
    replaced = _recolor_prs(prs, {old_hex: new_hex})
    if replaced:
        _save(prs, pptx_path)
    return replaced


def _theme_hex(theme: dict) -> dict[str, str]:
    return {
        key: f"#{str(value).upper()}"
        for key, value in theme.items()
        if key in {"bg", "bg_alt", "primary", "secondary", "accent", "text", "text_muted", "white", "dark"}
    }


def _observed_colors(prs: Presentation) -> Counter[str]:
    colors: Counter[str] = Counter()
    for slide in prs.slides:
        for shape in _walk_shapes(slide.shapes):
            fill = getattr(shape, "fill", None)
            if fill is not None:
                try:
                    color = _color_hex(fill.fore_color)
                    if color:
                        colors[color] += 1
                except (AttributeError, TypeError, ValueError):
                    pass
            line = getattr(shape, "line", None)
            if line is not None:
                try:
                    color = _color_hex(line.color)
                    if color:
                        colors[color] += 1
                except (AttributeError, TypeError, ValueError):
                    pass
            if getattr(shape, "has_text_frame", False):
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        color = _color_hex(run.font.color)
                        if color:
                            colors[color] += 1
    return colors


def swap_theme(pptx_path: str | Path, new_theme_key: str) -> dict:
    if new_theme_key not in THEMES:
        raise KeyError(f"Unknown theme: {new_theme_key}")
    prs = Presentation(str(pptx_path))
    observed = _observed_colors(prs)
    scores = {}
    for key, theme in THEMES.items():
        expected = set(_theme_hex(theme).values())
        scores[key] = sum(observed[color] for color in expected)
    old_key = max(scores, key=scores.get)
    if scores[old_key] < 2:
        return {
            "error": "无法可靠识别原主题；请先 inspect，再用 recolor 手动映射颜色。",
            "confidence": "low",
            "candidate": old_key,
        }
    old_theme = _theme_hex(THEMES[old_key])
    new_theme = _theme_hex(THEMES[new_theme_key])
    mapping = {old_theme[field]: new_theme[field] for field in old_theme if field in new_theme}
    replaced = _recolor_prs(prs, mapping)
    if replaced:
        _save(prs, pptx_path)
    return {
        "from_theme": old_key,
        "to_theme": new_theme_key,
        "replaced": replaced,
        "guessed": True,
        "confidence": "high" if scores[old_key] >= 6 else "medium",
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Edit an existing PPTX in place")
    sub = parser.add_subparsers(dest="command", required=True)
    text = sub.add_parser("text")
    text.add_argument("pptx")
    text.add_argument("--slide", type=int, required=True)
    text.add_argument("--find", required=True)
    text.add_argument("--replace", required=True)
    role = sub.add_parser("role")
    role.add_argument("pptx")
    role.add_argument("--slide", type=int, required=True)
    role.add_argument("--role", required=True)
    role.add_argument("--text", required=True)
    image = sub.add_parser("image")
    image.add_argument("pptx")
    image.add_argument("--slide", type=int, required=True)
    image.add_argument("--image", type=int, required=True)
    image.add_argument("--path", required=True)
    color = sub.add_parser("recolor")
    color.add_argument("pptx")
    color.add_argument("--old", required=True)
    color.add_argument("--new", required=True)
    theme = sub.add_parser("theme")
    theme.add_argument("pptx")
    theme.add_argument("--new", required=True)
    args = parser.parse_args()
    if args.command == "text":
        result = edit_text(args.pptx, args.slide, args.find, args.replace)
    elif args.command == "role":
        result = edit_text_by_role(args.pptx, args.slide, args.role, args.text)
    elif args.command == "image":
        result = swap_image(args.pptx, args.slide, args.image, args.path)
    elif args.command == "recolor":
        result = recolor(args.pptx, args.old, args.new)
    else:
        result = swap_theme(args.pptx, args.new)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
