"""Insert, delete, move, duplicate, and redraw PPTX slides."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from ppt_project import create_backup
from pptx import Presentation
from pptx_helper import (
    LAYOUT_REGISTRY,
    Section,
    _auto_search_images,
    choose_layout,
    choose_theme,
)
from reference_ppt import (
    _body_candidates,
    _remove_slide,
    _title_candidate,
    clone_slide,
    extract_template_profile,
)

from pptx_skill._io import save_prs as _save_prs_impl


def _position(index: int, length: int, allow_end: bool = False) -> int:
    """Convert a 1-based *index* to a 0-based position for internal use."""
    position = index - 1
    upper = length if allow_end else length - 1
    if not 0 <= position <= upper:
        raise IndexError(f"slide_index {index} out of range (1..{length})")
    return position


def _load_section_arg(value: str | Path) -> dict:
    """从 CLI 参数加载章节定义。

    支持两种形式：
    - 指向 JSON 文件的路径
    - 内联的 JSON 字符串

    统一处理 JSONDecodeError 抛出错信息提示用户，而不是裸异常崩溃。
    """
    section_path = Path(value)
    try:
        if section_path.exists():
            return json.loads(section_path.read_text(encoding="utf-8"))
        return json.loads(value)
    except json.JSONDecodeError as exc:
        raise ValueError(
            f"--section 参数既不是有效文件也不是合法 JSON: {exc}"
        ) from exc


def _move_id(prs: Presentation, old_position: int, new_position: int) -> None:
    slide_ids = prs.slides._sldIdLst
    element = slide_ids[old_position]
    slide_ids.remove(element)
    slide_ids.insert(new_position, element)


def _theme_for_existing(pptx_path: str | Path, theme_key: str | None) -> dict:
    if theme_key:
        return choose_theme(theme_key)
    try:
        return choose_theme(profile=extract_template_profile(pptx_path))
    except Exception:
        return choose_theme("editorial")


def _section_context(section: Section) -> dict:
    return {
        "title": section.title,
        "subtitle": section.subtitle,
        "bullets": section.bullets,
        "images": section.images,
        "kicker": section.kicker,
        "section_number": section.section_number,
        "page_number": section.page_number,
        "metrics": section.metrics,
        "events": section.events,
        "steps": section.steps,
        "table_headers": section.table_headers,
        "table_rows": section.table_rows,
        "left": section.left,
        "right": section.right,
        "quote": section.quote,
        "source": section.source,
        "layout_opts": section.layout_opts,
    }


def _insert_into_prs(
    prs: Presentation,
    section_dict: dict,
    position: int,
    theme: dict,
    layout: str | None,
    image_dir: str | Path,
    auto_search_images: bool,
) -> str:
    valid = set(Section.__dataclass_fields__)
    section = Section(**{key: value for key, value in section_dict.items() if key in valid})
    section.page_number = position + 1
    section.section_number = section.section_number or str(position).zfill(2)
    if auto_search_images and section.image_query and not section.images:
        _auto_search_images([section], str(image_dir), "zh")
    resolved_layout = layout or section.layout or choose_layout(section, 2, 10)
    if resolved_layout not in LAYOUT_REGISTRY:
        raise KeyError(f"Unknown layout: {resolved_layout}")
    LAYOUT_REGISTRY[resolved_layout](prs, theme, _section_context(section))
    _move_id(prs, len(prs.slides) - 1, position)
    return resolved_layout


def insert_slide(
    pptx_path: str | Path,
    index: int,
    section_dict: dict,
    layout: str | None = None,
    theme_key: str | None = None,
    auto_search_images: bool = True,
) -> int:
    path = Path(pptx_path).resolve()
    prs = Presentation(str(path))
    position = _position(index, len(prs.slides), allow_end=True)
    theme = _theme_for_existing(path, theme_key)
    image_dir = path.parent / "ppt_images"
    # 显式创建图片目录，避免 auto_search_images 关闭时下游访问不存在的目录。
    image_dir.mkdir(exist_ok=True)
    _insert_into_prs(
        prs, section_dict, position, theme, layout,
        image_dir, auto_search_images,
    )
    create_backup(path)
    _save_prs_impl(prs, path, backup=False)
    return len(prs.slides)


def delete_slide(pptx_path: str | Path, index: int) -> int:
    path = Path(pptx_path).resolve()
    prs = Presentation(str(path))
    position = _position(index, len(prs.slides))
    _remove_slide(prs, prs.slides[position])
    create_backup(path)
    _save_prs_impl(prs, path, backup=False)
    return len(prs.slides)


def move_slide(pptx_path: str | Path, from_index: int, to_index: int) -> None:
    path = Path(pptx_path).resolve()
    prs = Presentation(str(path))
    old = _position(from_index, len(prs.slides))
    new = _position(to_index, len(prs.slides))
    _move_id(prs, old, new)
    create_backup(path)
    _save_prs_impl(prs, path, backup=False)


def duplicate_slide(pptx_path: str | Path, source_index: int, target_index: int) -> int:
    path = Path(pptx_path).resolve()
    prs = Presentation(str(path))
    source_position = _position(source_index, len(prs.slides))
    target_position = _position(target_index, len(prs.slides), allow_end=True)
    clone_slide(prs, prs.slides[source_position])
    _move_id(prs, len(prs.slides) - 1, target_position)
    create_backup(path)
    _save_prs_impl(prs, path, backup=False)
    return len(prs.slides)


def _extract_basic_section(slide: Any) -> dict:
    title_shape = _title_candidate(slide)
    body_shapes = _body_candidates(slide, title_shape)
    bullets = []
    if body_shapes:
        bullets = [line.strip().lstrip("•").strip() for line in body_shapes[0].text.splitlines() if line.strip()]
    return {
        "title": title_shape.text.strip() if title_shape is not None else "",
        "bullets": bullets,
    }


def replace_layout(
    pptx_path: str | Path,
    index: int,
    new_layout: str,
    section_dict: dict | None = None,
    theme_key: str | None = None,
) -> None:
    path = Path(pptx_path).resolve()
    prs = Presentation(str(path))
    position = _position(index, len(prs.slides))
    old_slide = prs.slides[position]
    content = section_dict if section_dict else _extract_basic_section(old_slide)
    theme = _theme_for_existing(path, theme_key)
    image_dir = path.parent / "ppt_images"
    image_dir.mkdir(exist_ok=True)
    _insert_into_prs(
        prs, content, position, theme, new_layout,
        image_dir, False,
    )
    _remove_slide(prs, old_slide)
    create_backup(path)
    _save_prs_impl(prs, path, backup=False)


def main() -> int:
    parser = argparse.ArgumentParser(description="Manage pages in an existing PPTX")
    sub = parser.add_subparsers(dest="command", required=True)
    insert = sub.add_parser("insert")
    insert.add_argument("pptx")
    insert.add_argument("--index", type=int, required=True)
    insert.add_argument("--section", required=True, help="JSON object or JSON file")
    insert.add_argument("--layout")
    insert.add_argument("--theme")
    insert.add_argument("--no-image-search", action="store_true")
    delete = sub.add_parser("delete")
    delete.add_argument("pptx")
    delete.add_argument("--index", type=int, required=True)
    move = sub.add_parser("move")
    move.add_argument("pptx")
    move.add_argument("--from", dest="from_index", type=int, required=True)
    move.add_argument("--to", dest="to_index", type=int, required=True)
    duplicate = sub.add_parser("duplicate")
    duplicate.add_argument("pptx")
    duplicate.add_argument("--source", type=int, required=True)
    duplicate.add_argument("--target", type=int, required=True)
    replace = sub.add_parser("replace-layout")
    replace.add_argument("pptx")
    replace.add_argument("--index", type=int, required=True)
    replace.add_argument("--layout", required=True)
    replace.add_argument("--section")
    replace.add_argument("--theme")
    args = parser.parse_args()
    if args.command == "insert":
        section = _load_section_arg(args.section)
        result = insert_slide(
            args.pptx, args.index, section, args.layout, args.theme,
            not args.no_image_search,
        )
    elif args.command == "delete":
        result = delete_slide(args.pptx, args.index)
    elif args.command == "move":
        move_slide(args.pptx, args.from_index, args.to_index)
        result = True
    elif args.command == "duplicate":
        result = duplicate_slide(args.pptx, args.source, args.target)
    else:
        section = _load_section_arg(args.section) if args.section else None
        replace_layout(args.pptx, args.index, args.layout, section, args.theme)
        result = True
    print(json.dumps(result, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
