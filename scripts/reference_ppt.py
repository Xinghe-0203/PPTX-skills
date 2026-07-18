"""Analyze, clone, and populate user-provided PPTX reference decks."""

from __future__ import annotations

import argparse
import colorsys
import copy
import json
import re
from collections import Counter
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


REL_ATTRS = {qn("r:embed"), qn("r:id"), qn("r:link")}
CONTENT_PLACEHOLDERS = {
    PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.SUBTITLE,
    PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.PICTURE,
    PP_PLACEHOLDER.CHART, PP_PLACEHOLDER.TABLE,
}
TITLE_PLACEHOLDERS = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
BODY_PLACEHOLDERS = {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT}


def _enum_name(value: Any) -> str:
    return getattr(value, "name", str(value))


def _rgb_hex(value: Any) -> str | None:
    try:
        rgb = value.rgb
    except (AttributeError, TypeError, ValueError):
        return None
    if rgb is None:
        return None
    text = str(rgb).strip().upper()
    if re.fullmatch(r"[0-9A-F]{6}", text):
        return f"#{text}"
    try:
        int_value = int(rgb)
    except (TypeError, ValueError):
        return None
    if 0 <= int_value <= 0xFFFFFF:
        return f"#{int_value:06X}"
    return None


def _fill_hex(shape: Any) -> str | None:
    try:
        if shape.fill.type is None:
            return None
        return _rgb_hex(shape.fill.fore_color)
    except (AttributeError, TypeError, ValueError):
        return None


def _line_hex(shape: Any) -> str | None:
    try:
        return _rgb_hex(shape.line.color)
    except (AttributeError, TypeError, ValueError):
        return None


def _font_snapshot(run: Any | None) -> dict:
    if run is None:
        return {}
    font = run.font
    return {
        "name": font.name,
        "size": font.size.pt if font.size else None,
        "bold": font.bold,
        "italic": font.italic,
        "color": _rgb_hex(font.color),
    }


def _shape_font_data(shape: Any) -> tuple[list[str], list[float], list[str]]:
    fonts: list[str] = []
    sizes: list[float] = []
    colors: list[str] = []
    if not getattr(shape, "has_text_frame", False):
        return fonts, sizes, colors
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.font.name:
                fonts.append(run.font.name)
            if run.font.size:
                sizes.append(round(run.font.size.pt, 2))
            color = _rgb_hex(run.font.color)
            if color:
                colors.append(color)
    return fonts, sizes, colors


def _shape_record(shape: Any, slide_w: int, slide_h: int, index: int) -> dict:
    fonts, sizes, font_colors = _shape_font_data(shape)
    placeholder_type = None
    placeholder_idx = None
    if shape.is_placeholder:
        placeholder_type = _enum_name(shape.placeholder_format.type)
        placeholder_idx = shape.placeholder_format.idx
    text = shape.text.strip() if getattr(shape, "has_text_frame", False) else ""
    return {
        "index": index,
        "shape_id": shape.shape_id,
        "name": shape.name,
        "shape_type": _enum_name(shape.shape_type),
        "placeholder_type": placeholder_type,
        "placeholder_idx": placeholder_idx,
        "x": round(shape.left / slide_w, 5),
        "y": round(shape.top / slide_h, 5),
        "w": round(shape.width / slide_w, 5),
        "h": round(shape.height / slide_h, 5),
        "text": text,
        "fill": _fill_hex(shape),
        "line": _line_hex(shape),
        "fonts": sorted(set(fonts)),
        "font_sizes": sorted(set(sizes)),
        "font_colors": sorted(set(font_colors)),
        "has_chart": bool(getattr(shape, "has_chart", False)),
        "has_table": bool(getattr(shape, "has_table", False)),
    }


def _classify_slide(records: list[dict], index: int, total: int) -> str:
    text = " ".join(record["text"] for record in records if record["text"]).casefold()
    pictures = sum(record["shape_type"] == "PICTURE" for record in records)
    charts = sum(record["has_chart"] for record in records)
    tables = sum(record["has_table"] for record in records)
    title_shapes = [r for r in records if r["placeholder_type"] in {"TITLE", "CENTER_TITLE"}]
    body_shapes = [r for r in records if r["placeholder_type"] in {"BODY", "OBJECT"}]
    text_shapes = [r for r in records if r["text"]]
    if index == 1:
        return "cover"
    if index == total and any(token in text for token in ("thank", "谢谢", "致谢", "q&a")):
        return "end"
    if charts or tables or len(re.findall(r"\d+(?:\.\d+)?%", text)) >= 2:
        return "dashboard"
    if pictures >= 3:
        return "image_grid"
    if len(body_shapes) >= 2:
        return "comparison"
    if any(token in text for token in ("timeline", "时间线", "里程碑")):
        return "timeline"
    if any(token in text for token in ("process", "流程", "步骤", "路径")):
        return "process"
    if len(text_shapes) <= 2 and title_shapes:
        return "section"
    if pictures:
        return "text_image" if len(text_shapes) > 1 else "full_image"
    return "bullets"


def _layout_record(layout: Any, slide_w: int, slide_h: int, index: int) -> dict:
    placeholders = []
    for placeholder in layout.placeholders:
        placeholders.append({
            "idx": placeholder.placeholder_format.idx,
            "type": _enum_name(placeholder.placeholder_format.type),
            "name": placeholder.name,
            "x": round(placeholder.left / slide_w, 5),
            "y": round(placeholder.top / slide_h, 5),
            "w": round(placeholder.width / slide_w, 5),
            "h": round(placeholder.height / slide_h, 5),
        })
    return {"index": index, "name": layout.name, "placeholders": placeholders}


def analyze_presentation(pptx_path: str | Path, output_json: str | Path | None = None) -> dict:
    path = Path(pptx_path).resolve()
    prs = Presentation(str(path))
    slide_w, slide_h = prs.slide_width, prs.slide_height
    slides = []
    color_counter: Counter[str] = Counter()
    font_counter: Counter[str] = Counter()
    slide_level_shape_counts: list[int] = []
    for slide_index, slide in enumerate(prs.slides, start=1):
        records = [
            _shape_record(shape, slide_w, slide_h, shape_index)
            for shape_index, shape in enumerate(slide.shapes, start=1)
        ]
        for record in records:
            for color in (record["fill"], record["line"], *record["font_colors"]):
                if color:
                    color_counter[color] += 1
            font_counter.update(record["fonts"])
        slide_level_shape_counts.append(sum(not shape.is_placeholder for shape in slide.shapes))
        slides.append({
            "index": slide_index,
            "layout": slide.slide_layout.name,
            "role": _classify_slide(records, slide_index, len(prs.slides)),
            "shape_count": len(records),
            "shapes": records,
        })

    layouts = [
        _layout_record(layout, slide_w, slide_h, index)
        for index, layout in enumerate(prs.slide_layouts)
    ]
    useful_layouts = sum(
        any(placeholder["type"] in {"TITLE", "CENTER_TITLE", "BODY", "OBJECT", "PICTURE"}
            for placeholder in layout["placeholders"])
        for layout in layouts
    )
    avg_slide_shapes = (
        sum(slide_level_shape_counts) / len(slide_level_shape_counts)
        if slide_level_shape_counts else 0
    )
    if useful_layouts >= 4 and avg_slide_shapes < 3:
        mode = "native"
    elif slides:
        mode = "clone"
    else:
        mode = "visual-rebuild"

    result = {
        "source": str(path),
        "slide_size": {
            "width_emu": slide_w,
            "height_emu": slide_h,
            "width_in": round(slide_w / 914400, 3),
            "height_in": round(slide_h / 914400, 3),
            "aspect_ratio": round(slide_w / slide_h, 4),
        },
        "slide_count": len(slides),
        "master_count": len(prs.slide_masters),
        "layouts": layouts,
        "slides": slides,
        "design_tokens": {
            "colors": color_counter.most_common(16),
            "fonts": font_counter.most_common(10),
        },
        "recommended_mode": mode,
        "mode_reason": {
            "useful_layouts": useful_layouts,
            "average_slide_level_shapes": round(avg_slide_shapes, 2),
        },
    }
    if output_json:
        target = Path(output_json)
        target.parent.mkdir(parents=True, exist_ok=True)
        with target.open("w", encoding="utf-8", newline="\n") as handle:
            json.dump(result, handle, ensure_ascii=False, indent=2)
            handle.write("\n")
    return result


def _color_saturation(color: str) -> float:
    rgb = tuple(int(color[i:i + 2], 16) / 255 for i in (1, 3, 5))
    return colorsys.rgb_to_hsv(*rgb)[1]


def _color_luminance(color: str) -> float:
    r, g, b = (int(color[i:i + 2], 16) for i in (1, 3, 5))
    return (0.299 * r + 0.587 * g + 0.114 * b) / 255


def extract_template_profile(pptx_path: str | Path, name: str | None = None) -> dict:
    analysis = analyze_presentation(pptx_path)
    colors = [item[0] for item in analysis["design_tokens"]["colors"]]
    colors = [color for color in colors if re.fullmatch(r"#[0-9A-F]{6}", color)]
    bg = next((c for c in colors if _color_luminance(c) > 0.88), "#FAFAF8")
    dark = next((c for c in colors if _color_luminance(c) < 0.22), "#1A1A1A")
    primary = next(
        (c for c in colors if _color_saturation(c) > 0.25 and 0.16 < _color_luminance(c) < 0.72),
        dark,
    )
    accent = next(
        (c for c in sorted(colors, key=_color_saturation, reverse=True) if c != primary),
        "#D9623B",
    )
    on_dark = _color_luminance(bg) < 0.42
    profile_name = name or Path(pptx_path).stem
    return {
        "id": re.sub(r"[^a-z0-9]+", "-", profile_name.lower()).strip("-") or "reference-template",
        "name": profile_name,
        "description": f"Extracted from {Path(pptx_path).name}",
        "theme": {
            "bg": bg,
            "bg_alt": colors[1] if len(colors) > 1 else bg,
            "primary": primary,
            "secondary": colors[2] if len(colors) > 2 else primary,
            "accent": accent,
            "text": "#F2EEE6" if on_dark else dark,
            "text_muted": "#A3AAA5" if on_dark else "#727874",
            "white": "#FFFFFF",
            "dark": dark,
            "on_dark": on_dark,
        },
        "fonts": {
            "primary": analysis["design_tokens"]["fonts"][0][0]
            if analysis["design_tokens"]["fonts"] else "Microsoft YaHei"
        },
        "reference": {
            "path": str(Path(pptx_path).resolve()),
            "recommended_mode": analysis["recommended_mode"],
            "slide_roles": {str(s["index"]): s["role"] for s in analysis["slides"]},
        },
    }


def _copy_relationships(source_slide: Any, destination_slide: Any) -> dict[str, str]:
    mapping: dict[str, str] = {}
    for rel in source_slide.part.rels.values():
        if rel.reltype.endswith("/slideLayout") or rel.reltype.endswith("/notesSlide"):
            continue
        if rel.is_external:
            new_rid = destination_slide.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
        else:
            new_rid = destination_slide.part.rels.get_or_add(rel.reltype, rel.target_part)
        mapping[rel.rId] = new_rid
    return mapping


def _remap_relationship_ids(element: Any, mapping: dict[str, str]) -> None:
    for node in element.iter():
        for attribute in REL_ATTRS:
            old_rid = node.get(attribute)
            if old_rid in mapping:
                node.set(attribute, mapping[old_rid])


def clone_slide(prs: Presentation, source_slide: Any) -> Any:
    """Clone a slide inside the same package, preserving design and media relationships."""
    destination = prs.slides.add_slide(source_slide.slide_layout)
    relationship_map = _copy_relationships(source_slide, destination)
    for shape in list(destination.shapes):
        destination.shapes._spTree.remove(shape._element)
    for source_shape in source_slide.shapes:
        cloned_shape = copy.deepcopy(source_shape._element)
        _remap_relationship_ids(cloned_shape, relationship_map)
        destination.shapes._spTree.insert_element_before(cloned_shape, "p:extLst")
    source_bg = source_slide._element.cSld.find(qn("p:bg"))
    destination_bg = destination._element.cSld.find(qn("p:bg"))
    if destination_bg is not None:
        destination._element.cSld.remove(destination_bg)
    if source_bg is not None:
        cloned_bg = copy.deepcopy(source_bg)
        _remap_relationship_ids(cloned_bg, relationship_map)
        destination._element.cSld.insert(0, cloned_bg)
    source_map = source_slide._element.find(qn("p:clrMapOvr"))
    destination_map = destination._element.find(qn("p:clrMapOvr"))
    if source_map is not None:
        cloned_map = copy.deepcopy(source_map)
        if destination_map is not None:
            destination._element.replace(destination_map, cloned_map)
        else:
            destination._element.append(cloned_map)
    return destination


def _remove_slide(prs: Presentation, slide: Any) -> None:
    slide_id_list = prs.slides._sldIdLst
    for slide_id in list(slide_id_list):
        if prs.part.related_part(slide_id.rId) is slide.part:
            prs.part.drop_rel(slide_id.rId)
            slide_id_list.remove(slide_id)
            return


def _find_shape(slide: Any, reference: str | int) -> Any:
    if isinstance(reference, int) or str(reference).isdigit():
        index = int(reference)
        for shape in slide.shapes:
            if shape.shape_id == index:
                return shape
        if 1 <= index <= len(slide.shapes):
            return slide.shapes[index - 1]
    for shape in slide.shapes:
        if shape.name == str(reference):
            return shape
    raise KeyError(f"Shape not found on slide: {reference}")


def _first_run(shape: Any) -> Any | None:
    if not getattr(shape, "has_text_frame", False):
        return None
    for paragraph in shape.text_frame.paragraphs:
        if paragraph.runs:
            return paragraph.runs[0]
    return None


def _apply_font_snapshot(run: Any, snapshot: dict) -> None:
    if snapshot.get("name"):
        run.font.name = snapshot["name"]
    if snapshot.get("size"):
        run.font.size = Pt(snapshot["size"])
    if snapshot.get("bold") is not None:
        run.font.bold = snapshot["bold"]
    if snapshot.get("italic") is not None:
        run.font.italic = snapshot["italic"]
    if snapshot.get("color"):
        color = snapshot["color"].lstrip("#")
        run.font.color.rgb = RGBColor.from_string(color)


def replace_text(shape: Any, text: str | None = None, bullets: list[str] | None = None) -> None:
    if not getattr(shape, "has_text_frame", False):
        raise TypeError(f"Shape '{shape.name}' does not contain text")
    snapshot = _font_snapshot(_first_run(shape))
    text_frame = shape.text_frame
    text_frame.clear()
    values = bullets if bullets is not None else str(text or "").splitlines()
    if not values:
        values = [""]
    for index, value in enumerate(values):
        paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
        paragraph.level = 0
        run = paragraph.add_run()
        run.text = f"• {value}" if bullets is not None and not str(value).lstrip().startswith("•") else str(value)
        _apply_font_snapshot(run, snapshot)


def replace_image(slide: Any, shape: Any, image_path: str | Path) -> Any:
    image_path = str(Path(image_path).resolve())
    picture = slide.shapes.add_picture(image_path, shape.left, shape.top, shape.width, shape.height)
    tree = shape._element.getparent()
    tree.remove(picture._element)
    shape._element.addprevious(picture._element)
    tree.remove(shape._element)
    return picture


def _delete_shape(shape: Any) -> None:
    parent = shape._element.getparent()
    if parent is not None:
        parent.remove(shape._element)


def apply_replacements(slide: Any, replacements: dict[str, Any]) -> set[str]:
    touched: set[str] = set()
    for reference, replacement in replacements.items():
        shape = _find_shape(slide, reference)
        touched.add(shape.name)
        if isinstance(replacement, str):
            replace_text(shape, text=replacement)
            continue
        if not isinstance(replacement, dict):
            replace_text(shape, text=str(replacement))
            continue
        if replacement.get("delete"):
            _delete_shape(shape)
        elif "image" in replacement:
            replace_image(slide, shape, replacement["image"])
        elif "bullets" in replacement:
            replace_text(shape, bullets=[str(item) for item in replacement["bullets"]])
        else:
            replace_text(shape, text=str(replacement.get("text", "")))
    return touched


def _content_lines(content: dict) -> list[str]:
    if content.get("bullets"):
        return [str(item) for item in content["bullets"]]
    if content.get("metrics"):
        return [f"{item.get('label', '')}: {item.get('value', '')} {item.get('change', '')}".strip()
                for item in content["metrics"]]
    if content.get("events"):
        return [f"{item.get('date', '')}  {item.get('title', '')}".strip() for item in content["events"]]
    if content.get("steps"):
        return [str(item) for item in content["steps"]]
    if content.get("left") or content.get("right"):
        lines = []
        for side in (content.get("left", {}), content.get("right", {})):
            if side:
                lines.append(str(side.get("title", "")))
                lines.extend(str(item) for item in side.get("bullets", []))
        return lines
    if content.get("quote"):
        return [str(content["quote"]), str(content.get("source", ""))]
    return []


def _title_candidate(slide: Any) -> Any | None:
    for shape in slide.shapes:
        if shape.is_placeholder and shape.placeholder_format.type in TITLE_PLACEHOLDERS:
            return shape
    candidates = [
        shape for shape in slide.shapes
        if getattr(shape, "has_text_frame", False) and shape.text.strip()
    ]
    if not candidates:
        return None
    return min(candidates, key=lambda shape: (shape.top, -shape.width))


def _body_candidates(slide: Any, title_shape: Any | None) -> list[Any]:
    title_id = title_shape.shape_id if title_shape is not None else None
    placeholders = [
        shape for shape in slide.shapes
        if shape.shape_id != title_id and shape.is_placeholder
        and shape.placeholder_format.type in BODY_PLACEHOLDERS
    ]
    if placeholders:
        return sorted(placeholders, key=lambda shape: shape.top)
    return sorted(
        [shape for shape in slide.shapes
         if shape.shape_id != title_id and getattr(shape, "has_text_frame", False)
         and shape.text.strip() and shape.width * shape.height > 0],
        key=lambda shape: shape.width * shape.height,
        reverse=True,
    )


def auto_bind_slide(slide: Any, content: dict) -> set[str]:
    touched: set[str] = set()
    title_shape = _title_candidate(slide)
    if title_shape is not None and content.get("title") is not None:
        replace_text(title_shape, text=str(content.get("title", "")))
        touched.add(title_shape.name)
    body_shapes = _body_candidates(slide, title_shape)
    if content.get("subtitle") and body_shapes:
        replace_text(body_shapes[0], text=str(content["subtitle"]))
        touched.add(body_shapes[0].name)
        body_shapes = body_shapes[1:]
    lines = _content_lines(content)
    if lines and body_shapes:
        replace_text(body_shapes[0], bullets=lines if content.get("bullets") else None,
                     text="\n".join(lines))
        touched.add(body_shapes[0].name)
    images = [
        Path(str(path)) for path in content.get("images", [])
        if path is not None and str(path).strip()
    ]
    pictures = [shape for shape in list(slide.shapes) if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    for picture, image in zip(pictures, images):
        if image.exists():
            touched.add(picture.name)
            replace_image(slide, picture, image)
    return touched


def _clear_unmapped_text(slide: Any, touched: set[str]) -> None:
    for shape in slide.shapes:
        if shape.name in touched or not getattr(shape, "has_text_frame", False):
            continue
        text = shape.text.strip()
        if not text:
            continue
        if re.fullmatch(r"\d+", text) and shape.height < Inches(0.45):
            continue
        _, font_sizes, _ = _shape_font_data(shape)
        likely_content = bool(font_sizes and max(font_sizes) >= 14 and len(text) > 2)
        if ((shape.is_placeholder and shape.placeholder_format.type in CONTENT_PLACEHOLDERS)
                or likely_content):
            replace_text(shape, text="")


def compose_from_reference(
    reference_pptx: str | Path,
    plan: dict | str | Path,
    output_path: str | Path,
) -> str:
    """Clone exemplar slides and apply exact shape-level replacements from a plan."""
    if isinstance(plan, (str, Path)):
        with Path(plan).open("r", encoding="utf-8") as handle:
            plan = json.load(handle)
    prs = Presentation(str(reference_pptx))
    source_slides = list(prs.slides)
    if not source_slides:
        raise ValueError("Reference presentation has no exemplar slides")
    for slide_plan in plan.get("slides", []):
        source_index = int(slide_plan.get("source_slide", 1))
        if not 1 <= source_index <= len(source_slides):
            raise IndexError(f"source_slide out of range: {source_index}")
        destination = clone_slide(prs, source_slides[source_index - 1])
        replacements = slide_plan.get("replacements", {})
        touched = apply_replacements(destination, replacements) if replacements else set()
        if slide_plan.get("content"):
            touched |= auto_bind_slide(destination, slide_plan["content"])
        if slide_plan.get("clear_unmapped_content", False):
            _clear_unmapped_text(destination, touched)
    if plan.get("remove_source_slides", True):
        for slide in source_slides:
            _remove_slide(prs, slide)
    target = Path(output_path).resolve()
    target.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(target))
    return str(target)


def _role_for_content(content: dict, position: int, total: int) -> str:
    if position == 0:
        return "cover"
    if position == total - 1:
        return "end"
    if content.get("metrics") or content.get("table_headers"):
        return "dashboard"
    if content.get("events"):
        return "timeline"
    if content.get("left") and content.get("right"):
        return "comparison"
    if content.get("steps"):
        return "process"
    if content.get("quote"):
        return "quote"
    if len(content.get("images", [])) >= 3:
        return "image_grid"
    if content.get("images"):
        return "text_image"
    return "bullets"


def _layout_score(layout: Any, role: str, content: dict) -> int:
    name = layout.name.casefold()
    placeholder_types = [placeholder.placeholder_format.type for placeholder in layout.placeholders]
    title_count = sum(item in TITLE_PLACEHOLDERS for item in placeholder_types)
    body_count = sum(item in BODY_PLACEHOLDERS for item in placeholder_types)
    picture_count = sum(item == PP_PLACEHOLDER.PICTURE for item in placeholder_types)
    score = title_count * 4
    role_tokens = {
        "cover": ("title", "cover", "封面", "标题"),
        "end": ("title", "closing", "结束", "致谢"),
        "comparison": ("comparison", "two", "对比", "两栏"),
        "dashboard": ("content", "chart", "data", "数据", "图表"),
        "timeline": ("timeline", "时间"),
        "process": ("process", "流程"),
        "text_image": ("picture", "image", "图片", "图文"),
        "image_grid": ("picture", "image", "图片", "图文"),
    }
    score += 8 * sum(token in name for token in role_tokens.get(role, ()))
    if role in {"cover", "end", "section"}:
        score += 6 if body_count == 0 else -body_count
    elif role == "comparison":
        score += body_count * 6
    else:
        score += body_count * 4
    if content.get("images"):
        score += picture_count * 8
    if content.get("subtitle"):
        score += sum(item == PP_PLACEHOLDER.SUBTITLE for item in placeholder_types) * 5
    return score


def _fill_native_slide(slide: Any, content: dict) -> None:
    title_shape = _title_candidate(slide)
    if title_shape is not None:
        replace_text(title_shape, text=str(content.get("title", "")))
    subtitle_shapes = [
        shape for shape in slide.placeholders
        if shape.placeholder_format.type == PP_PLACEHOLDER.SUBTITLE
    ]
    if subtitle_shapes and content.get("subtitle"):
        replace_text(subtitle_shapes[0], text=str(content["subtitle"]))
    body_shapes = [
        shape for shape in slide.placeholders
        if shape.placeholder_format.type in BODY_PLACEHOLDERS
    ]
    lines = _content_lines(content)
    if lines and body_shapes:
        replace_text(
            body_shapes[0],
            bullets=lines if content.get("bullets") else None,
            text="\n".join(lines),
        )
    elif lines:
        box = slide.shapes.add_textbox(Inches(0.9), Inches(1.65), Inches(11.5), Inches(4.9))
        replace_text(box, bullets=lines if content.get("bullets") else None, text="\n".join(lines))
        for paragraph in box.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(18)
    images = [Path(path) for path in content.get("images", []) if Path(path).exists()]
    picture_placeholders = [
        shape for shape in slide.placeholders
        if shape.placeholder_format.type == PP_PLACEHOLDER.PICTURE
    ]
    for placeholder, image in zip(picture_placeholders, images):
        placeholder.insert_picture(str(image.resolve()))


def generate_native_from_reference(
    reference_pptx: str | Path,
    title: str,
    sections: list[dict],
    output_path: str | Path,
    subtitle: str = "",
) -> str:
    """Create new slides from the reference deck's native masters and layouts."""
    prs = Presentation(str(reference_pptx))
    source_slides = list(prs.slides)
    for slide in source_slides:
        _remove_slide(prs, slide)
    items = [{"title": title, "subtitle": subtitle}]
    items.extend(sections)
    items.append({"title": "谢谢", "subtitle": subtitle})
    for position, content in enumerate(items):
        role = _role_for_content(content, position, len(items))
        layout = max(prs.slide_layouts, key=lambda item: _layout_score(item, role, content))
        slide = prs.slides.add_slide(layout)
        _fill_native_slide(slide, content)
    target = Path(output_path).resolve()
    target.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(target))
    return str(target)


def _choose_exemplar(analysis: dict, role: str, used: Counter[int]) -> int:
    matching = [slide["index"] for slide in analysis["slides"] if slide["role"] == role]
    candidates = matching or [slide["index"] for slide in analysis["slides"]]
    return min(candidates, key=lambda index: used[index])


def generate_from_reference(
    reference_pptx: str | Path,
    title: str,
    sections: list[dict],
    output_path: str | Path,
    subtitle: str = "",
    mode: str = "auto",
) -> str:
    """Reuse a reference deck in native-layout or exact exemplar-clone mode."""
    analysis = analyze_presentation(reference_pptx)
    resolved_mode = analysis["recommended_mode"] if mode == "auto" else mode
    if resolved_mode == "native":
        return generate_native_from_reference(
            reference_pptx, title, sections, output_path, subtitle,
        )
    if resolved_mode not in {"clone", "visual-rebuild"}:
        raise ValueError(f"Unsupported reference mode: {resolved_mode}")
    items = [{"title": title, "subtitle": subtitle}]
    items.extend(sections)
    items.append({"title": "谢谢", "subtitle": subtitle})
    usage: Counter[int] = Counter()
    slides = []
    for position, content in enumerate(items):
        role = _role_for_content(content, position, len(items))
        source_index = _choose_exemplar(analysis, role, usage)
        usage[source_index] += 1
        slides.append({
            "source_slide": source_index,
            "content": content,
            "clear_unmapped_content": True,
        })
    return compose_from_reference(reference_pptx, {"slides": slides}, output_path)


def _build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Analyze and reuse user-provided PPTX decks")
    sub = parser.add_subparsers(dest="command", required=True)

    analyze = sub.add_parser("analyze", help="Extract layouts, shape names, roles, colors and fonts")
    analyze.add_argument("pptx")
    analyze.add_argument("--output")

    profile = sub.add_parser("profile", help="Extract a reusable visual profile")
    profile.add_argument("pptx")
    profile.add_argument("--name")
    profile.add_argument("--output", required=True)

    compose = sub.add_parser("compose", help="Clone slides according to a shape-level plan")
    compose.add_argument("pptx")
    compose.add_argument("--plan", required=True)
    compose.add_argument("--output", required=True)

    generate = sub.add_parser("generate", help="Auto-bind structured content to exemplar slides")
    generate.add_argument("pptx")
    generate.add_argument("--content", required=True)
    generate.add_argument("--output", required=True)
    generate.add_argument("--mode", choices=("auto", "native", "clone", "visual-rebuild"), default="auto")
    return parser


def main() -> int:
    args = _build_parser().parse_args()
    if args.command == "analyze":
        result = analyze_presentation(args.pptx, args.output)
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0
    if args.command == "profile":
        result = extract_template_profile(args.pptx, args.name)
        target = Path(args.output)
        target.parent.mkdir(parents=True, exist_ok=True)
        with target.open("w", encoding="utf-8", newline="\n") as handle:
            json.dump(result, handle, ensure_ascii=False, indent=2)
            handle.write("\n")
        print(str(target.resolve()))
        return 0
    if args.command == "compose":
        print(compose_from_reference(args.pptx, args.plan, args.output))
        return 0

    with Path(args.content).open("r", encoding="utf-8") as handle:
        content = json.load(handle)
    print(generate_from_reference(
        args.pptx,
        content["title"],
        content.get("sections", []),
        args.output,
        content.get("subtitle", ""),
        args.mode,
    ))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
