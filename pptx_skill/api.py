"""Public API and compatibility facade for the PPTX skill.

This module preserves the historical function signatures of ``auto_generate_ppt``
and ``auto_validate_ppt`` while adding the new structured result types and
experimental switches for the V2 adaptive/QA pipeline.
"""
from __future__ import annotations

import os
from pathlib import Path
from typing import Any, Literal

from pptx_skill.content_adapter import adapt_legacy_sections
from pptx_skill.generation_pipeline import GenerationResult as PipelineGenerationResult
from pptx_skill.manifest import ManifestV3, load_manifest, save_manifest_v3, set_current_content
from pptx_skill.visual_qa import CheckOutcome, QACheckResult, QAIssue, QAReport, Severity

# Keep api.GenerationResult as the canonical public result type while sharing
# the same field layout with the pipeline's GenerationResult.
GenerationResult = PipelineGenerationResult

def _ensure_dir(path: str) -> None:
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)


def _legacy_validate(pptx_path: str) -> dict[str, Any]:
    """Port the existing structural validation to the new outcome model.

    Mirrors ``pptx_helper.auto_validate_ppt`` so callers switching to the
    facade get the same five checks (adjacent diversity, layout diversity,
    cover content, font-size hierarchy, color restraint) instead of the
    stripped-down three-check variant that previously lived here.
    """
    from pptx import Presentation

    if not os.path.exists(pptx_path):
        return {"passed": False, "checks": [], "warnings": [f"文件不存在: {pptx_path}"]}

    prs = Presentation(pptx_path)
    slides = list(prs.slides)
    n_slides = len(slides)
    checks: list[tuple[str, bool | None]] = []
    warnings: list[str] = []
    passed = True

    # 1. 相邻版式不重复（连续 3 页 shape 数量相同视为可能重复）。
    # ``i`` is 0-based over ``slides``; the three repeated pages are the
    # 1-based slides ``i-1, i, i+1``, so the range label ``第{i-1}-{i+1}页``
    # is already 1-based and correct.
    if n_slides >= 3:
        counts = [len(slide.shapes) for slide in slides]
        for i in range(2, n_slides):
            if counts[i] == counts[i - 1] == counts[i - 2]:
                warnings.append(
                    f"第{i-1}-{i+1}页shape数量相同({counts[i]})，可能版式重复"
                )
    checks.append(("相邻版式多样性", not any("版式重复" in w for w in warnings)))

    # 2. 版式多样性：shape 数量不重复率 > 50%。
    unique_counts = len({len(s.shapes) for s in slides})
    diversity_ratio = unique_counts / max(n_slides, 1)
    checks.append(("版式多样性", diversity_ratio >= 0.5))
    if diversity_ratio < 0.5:
        warnings.append(f"版式多样性不足: {unique_counts}/{n_slides}种不同的shape数量")
        passed = False

    # 3. 封面非空。
    if n_slides >= 1:
        cover_shapes = len(slides[0].shapes)
        checks.append(("封面有内容", cover_shapes >= 2))
        if cover_shapes < 2:
            warnings.append("封面内容过少")
            passed = False

    # 4. 字号层级跳跃（最大/最小 ≥ 2.5）。
    max_font = 0
    min_font = 999
    for slide in slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if run.font.size:
                        pt_val = run.font.size.pt
                        max_font = max(max_font, pt_val)
                        min_font = min(min_font, pt_val)
    if max_font > 0 and min_font < 999:
        ratio = max_font / max(min_font, 1)
        checks.append(("字号层级跳跃", ratio >= 2.5))
        if ratio < 2.5:
            warnings.append(f"字号层级不够跳跃: 最大{max_font}pt/最小{min_font}pt={ratio:.1f}x (需≥2.5x)")
            passed = False
    else:
        warnings.append("无法检测字号层级（可能未设置font.size）")

    # 5. 配色克制（≤ 10 种填充色）。
    colors: set[str] = set()
    for slide in slides:
        for shape in slide.shapes:
            try:
                if hasattr(shape, "fill") and shape.fill.type is not None:
                    if shape.fill.fore_color and shape.fill.fore_color.rgb:
                        colors.add(str(shape.fill.fore_color.rgb))
            except Exception:
                pass
    checks.append(("配色克制", len(colors) <= 10))
    if len(colors) > 10:
        warnings.append(f"使用颜色过多: {len(colors)}种 (建议≤10)")
        passed = False

    return {"passed": passed, "checks": checks, "warnings": warnings, "total_slides": n_slides}


def _qa_report_to_legacy(report: QAReport) -> dict[str, Any]:
    """Map a new QAReport back to the legacy validation dict shape."""
    blockers = [i for i in report.issues if i.severity == Severity.BLOCKER]
    warnings = [i for i in report.issues if i.severity == Severity.WARNING]
    return {
        "passed": report.status == CheckOutcome.PASS and not blockers,
        "checks": [(c.check_id, c.outcome == CheckOutcome.PASS) for c in report.checks],
        "warnings": [i.message for i in warnings + blockers],
        "total_slides": report.metrics.get("total_slides", 0),
        "_qa_report": report,
    }


def _build_qa_report(pptx_path: str, legacy_result: dict[str, Any] | None = None) -> QAReport:
    """Build a minimal QAReport from legacy structural validation.

    Real semantic/rendered checks are added in later PRs; this stub preserves
    the contract and lets callers switch to the new result type immediately.
    """
    if legacy_result is None:
        legacy_result = _legacy_validate(pptx_path)

    checks: list[QACheckResult] = []
    for name, ok in legacy_result.get("checks", []):
        checks.append(
            QACheckResult(
                check_id=name,
                outcome=CheckOutcome.PASS if ok else CheckOutcome.FAIL,
                confidence=1.0,
            )
        )

    issues = []
    for warning in legacy_result.get("warnings", []):
        issues.append(
            QAIssue(
                code="LEGACY_WARNING",
                severity=Severity.WARNING,
                slide_index=-1,
                message=warning,
            )
        )

    status = CheckOutcome.PASS if legacy_result.get("passed") else CheckOutcome.FAIL
    return QAReport(
        status=status,
        checks=checks,
        issues=issues,
        metrics={"total_slides": legacy_result.get("total_slides", 0)},
    )


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def auto_generate_ppt(
    title,
    subtitle: str = "",
    sections: list[dict[str, Any] | Any] | None = None,
    output_path: str = "output.pptx",
    theme_key: str | None = None,
    image_dir: str = "./ppt_images",
    auto_search_images: bool = True,
    lang: str = "zh",
    template_key: str | None = None,
    template_profile: dict[str, Any] | None = None,
    template_catalog_path: str | None = None,
    layout_engine: Literal["legacy", "adaptive"] = "legacy",
    qa_mode: Literal["off", "report", "strict"] = "off",
    auto_repair: bool = True,
    max_repair_passes: int = 2,
    renderer_engine: Literal["auto", "libreoffice", "com"] = "auto",
    target_dpi: int = 150,
    reference_baseline: str | None = None,
    return_result: bool = False,
) -> str | GenerationResult:
    """Generate a PPTX. Default behavior is identical to the legacy function.

    New optional arguments are stubs for the V2 pipeline:

    - ``layout_engine="adaptive"`` delegates to the constraint-based adaptive
      pipeline in ``generation_pipeline.run_generation_pipeline``.
    - ``qa_mode`` controls whether a QA report is produced; ``strict`` would
      block on blockers once real visual QA is wired in.
    - ``return_result=True`` returns a structured ``GenerationResult``.
    """
    if layout_engine not in ("legacy", "adaptive"):
        raise ValueError(f"Unknown layout_engine: {layout_engine}")

    sections = sections or []
    abs_output = os.path.abspath(output_path)
    _ensure_dir(abs_output)

    if layout_engine == "adaptive":
        from pptx_skill.generation_pipeline import run_generation_pipeline

        content = adapt_legacy_sections(title, subtitle, sections, locale=lang or "zh-CN")

        baseline_pngs = None
        if reference_baseline:
            baseline_dir = Path(reference_baseline)
            if baseline_dir.is_dir():
                baseline_pngs = sorted(str(p) for p in baseline_dir.glob("*.png"))

        result = run_generation_pipeline(
            content=content,
            output_path=abs_output,
            profile_state=template_profile,
            qa_mode=qa_mode,
            auto_repair=auto_repair,
            max_repair_passes=max_repair_passes,
            renderer_engine=renderer_engine,
            target_dpi=target_dpi,
            reference_baseline_pngs=baseline_pngs,
        )
        return result

    # Legacy generation path.
    from pptx_helper import auto_generate_ppt as legacy_auto_generate_ppt

    pptx_path = legacy_auto_generate_ppt(
        title=title,
        subtitle=subtitle,
        sections=sections,
        output_path=abs_output,
        theme_key=theme_key,
        image_dir=image_dir,
        auto_search_images=auto_search_images,
        lang=lang,
        template_key=template_key,
        template_profile=template_profile,
        template_catalog_path=template_catalog_path,
    )

    # Build and persist V3 content model (best-effort; never breaks legacy path).
    manifest = None
    try:
        content = adapt_legacy_sections(title, subtitle, sections, locale=lang or "zh-CN")
        manifest = load_manifest(pptx_path) or ManifestV3()
        set_current_content(manifest, content)
        manifest.legacy["theme_key"] = theme_key
        manifest.legacy["template_key"] = template_key
        save_manifest_v3(pptx_path, manifest)
    except Exception:
        # V3 write failed — the legacy path already wrote an accurate V2
        # manifest (content matches the generated deck), and load_manifest()
        # auto-migrates V2→V3 on read, so we leave it in place. Clearing the
        # embedded part would orphan its relationship in presentation.xml.rels
        # and trigger PowerPoint's repair prompt.
        import logging

        logging.getLogger(__name__).warning(
            "V3 manifest write failed; leaving accurate V2 manifest in place",
            exc_info=True,
        )

    qa_report: QAReport | None = None
    if qa_mode != "off":
        qa_report = _build_qa_report(pptx_path)
        if qa_mode == "strict" and qa_report.status != CheckOutcome.PASS:
            # In PR1 strict only blocks on the legacy structural checks.
            from pptx_skill.visual_qa import PresentationQualityError

            raise PresentationQualityError(qa_report)

    if not return_result:
        return pptx_path

    return GenerationResult(
        pptx_path=pptx_path,
        manifest_path=_sidecar_path(pptx_path),
        qa_report_path=None,
        qa_status=qa_report.status.value if qa_report else None,
        render_trace_path=None,
        preview_dir=None,
        repair_passes=0,
        manifest=manifest,
    )


def auto_validate_ppt(pptx_path: str, return_report: bool = False) -> dict[str, Any] | QAReport:
    """Validate a PPTX.

    By default returns the legacy dict for backward compatibility. Pass
    ``return_report=True`` to receive the new ``QAReport``.
    """
    legacy = _legacy_validate(pptx_path)
    report = _build_qa_report(pptx_path, legacy)
    if return_report:
        return report
    return _qa_report_to_legacy(report)


def _sidecar_path(pptx_path: str) -> str:
    return str(Path(pptx_path).with_suffix(".manifest.json"))
