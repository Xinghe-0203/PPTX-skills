"""Regression tests for SVG import path handling."""

from __future__ import annotations

import sys
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from PIL import Image
from pptx import Presentation

from pptx_skill.svg_import import import_svg_as_image


class SvgImportRegressionTests(unittest.TestCase):
    def test_raster_import_persists_when_presentation_is_a_path_object(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            deck_path = root / "deck.pptx"
            prs = Presentation()
            prs.slides.add_slide(prs.slide_layouts[6])
            prs.save(deck_path)

            svg_path = root / "shape.svg"
            svg_path.write_text(
                '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 72 72">'
                '<rect width="72" height="72"/></svg>',
                encoding="utf-8",
            )
            png = BytesIO()
            Image.new("RGB", (4, 4), "red").save(png, "PNG")
            fake_cairosvg = SimpleNamespace(svg2png=lambda **kwargs: png.getvalue())

            with patch.dict(sys.modules, {"cairosvg": fake_cairosvg}):
                shape_name = import_svg_as_image(
                    deck_path,
                    1,
                    svg_path=str(svg_path),
                    width=1,
                    height=1,
                    name="SVG_Path_Object",
                )

            reopened = Presentation(deck_path)
            self.assertEqual(shape_name, "SVG_Path_Object")
            self.assertEqual(len(reopened.slides[0].shapes), 1)
            self.assertEqual(reopened.slides[0].shapes[0].name, "SVG_Path_Object")


if __name__ == "__main__":
    unittest.main()
