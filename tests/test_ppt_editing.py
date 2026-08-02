import sys
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS))

from ppt_edit import edit_text, recolor  # noqa: E402
from ppt_pages import delete_slide, duplicate_slide, insert_slide, move_slide, replace_layout  # noqa: E402
from ppt_project import load_project  # noqa: E402
from pptx_helper import auto_generate_ppt  # noqa: E402


def deck_text(path: Path) -> str:
    prs = Presentation(str(path))
    return "\n".join(
        shape.text for slide in prs.slides for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
    )


class PptEditingTests(unittest.TestCase):
    def _make_deck(self, directory: Path) -> Path:
        path = directory / "deck.pptx"
        auto_generate_ppt(
            title="原始标题",
            subtitle="副标题",
            sections=[
                {"title": "第一页", "bullets": ["要点 A", "要点 B"]},
                {"title": "第二页", "steps": ["发现", "设计", "交付"]},
            ],
            output_path=str(path),
            theme_key="editorial",
            auto_search_images=False,
        )
        return path

    def test_manifest_is_embedded_and_sidecar_fallback_exists(self):
        with tempfile.TemporaryDirectory() as tmp:
            directory = Path(tmp)
            deck = self._make_deck(directory)
            sidecar = deck.with_suffix(".manifest.json")
            self.assertTrue(sidecar.exists())
            project = load_project(deck)
            self.assertEqual(project["title"], "原始标题")
            sidecar.unlink()
            embedded = load_project(deck)
            self.assertEqual(embedded["sections"][0]["title"], "第一页")
            Presentation(str(deck))

    def test_facade_generated_deck_is_round_trippable(self):
        # Decks produced by the V3 facade (pptx_skill.auto_generate_ppt) carry a
        # Manifest V3, not a v2 sections payload. load_project must still rebuild
        # a v2 project so edit_section/regenerate keep working.
        from pptx_skill import auto_generate_ppt as facade_generate
        from ppt_project import edit_section, regenerate

        with tempfile.TemporaryDirectory() as tmp:
            directory = Path(tmp)
            deck = directory / "facade.pptx"
            facade_generate(
                title="封面标题",
                subtitle="副标题",
                sections=[
                    {"title": "第一页", "bullets": ["要点 A", "要点 B"]},
                    {"title": "指标页", "metrics": [{"label": "x", "value": "10", "change": "+5%"}]},
                ],
                output_path=str(deck),
                auto_search_images=False,
            )
            project = load_project(deck)
            self.assertIsNotNone(project)
            self.assertEqual(len(project["sections"]), 2)
            self.assertEqual(project["sections"][0]["title"], "第一页")
            self.assertEqual(project["sections"][0]["layout"], "bullets")
            self.assertEqual(project["sections"][1]["layout"], "dashboard")

            # edit_section must succeed on a V3-only deck.
            out = edit_section(deck, 1, {"title": "改后标题"})
            prs = Presentation(out)
            text = "\n".join(
                shape.text for slide in prs.slides for shape in slide.shapes
                if getattr(shape, "has_text_frame", False)
            )
            self.assertIn("改后标题", text)

            # regenerate to a new path must also succeed.
            other = directory / "regen.pptx"
            project["sections"][0]["title"] = "再次改"
            regenerate(project, other)
            self.assertTrue(other.exists())

    def test_regenerate_preserves_slide_count_with_toc_inserted(self):
        # When >3 sections auto_generate_ppt inserts a TOC page. The layouts
        # sequence (cover,toc,...,end) must stay aligned with sections after
        # the cover/toc/end are filtered out, so regenerate reproduces the
        # same slide count instead of silently truncating via zip().
        import copy
        import sys as _sys
        _sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
        from pptx_helper import auto_generate_ppt as legacy_generate
        from ppt_project import regenerate

        with tempfile.TemporaryDirectory() as tmp:
            directory = Path(tmp)
            deck = directory / "deck.pptx"
            legacy_generate(
                title="多页",
                sections=[
                    {"title": f"第{i}页", "bullets": ["a", "b"]}
                    for i in range(1, 5)
                ],
                output_path=str(deck),
                theme_key="editorial",
                auto_search_images=False,
            )
            project = load_project(deck)
            # V2 manifest stores the full used_layouts incl. cover/toc/end.
            self.assertIn("toc", project["layouts"])
            mutated = copy.deepcopy(project)
            mutated["sections"][2]["title"] = "改后"
            out = directory / "regen.pptx"
            regenerate(mutated, out)
            result = Presentation(out)
            # Original had cover + toc + 4 content + end = 7 slides.
            self.assertEqual(len(result.slides), 7)
            text = "\n".join(
                shape.text for slide in result.slides for shape in slide.shapes
                if getattr(shape, "has_text_frame", False)
            )
            self.assertIn("改后", text)

    def test_text_color_and_page_operations_create_backups(self):
        with tempfile.TemporaryDirectory() as tmp:
            directory = Path(tmp)
            deck = self._make_deck(directory)
            self.assertEqual(edit_text(deck, 1, "原始标题", "修改后标题"), 1)
            self.assertIn("修改后标题", deck_text(deck))
            self.assertTrue(directory.joinpath("deck.bak.pptx").exists())
            self.assertGreater(recolor(deck, "#1A1A1A", "#123456"), 0)

            self.assertEqual(duplicate_slide(deck, 2, 3), 5)
            self.assertEqual(delete_slide(deck, 3), 4)
            move_slide(deck, 2, 3)
            self.assertEqual(insert_slide(
                deck, 2,
                {"title": "插入页", "bullets": ["新增内容"]},
                layout="bullets", auto_search_images=False,
            ), 5)
            self.assertIn("插入页", deck_text(deck))
            replace_layout(
                deck, 2, "quote",
                {"title": "引言", "quote": "重新设计这一页", "source": "Test"},
            )
            self.assertIn("重新设计这一页", deck_text(deck))


if __name__ == "__main__":
    unittest.main()
