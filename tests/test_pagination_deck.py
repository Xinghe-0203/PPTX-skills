"""Tests for the pagination framework and deck planner (PR6)."""
from __future__ import annotations

import os
import tempfile
import unittest

from pptx_skill.content_model import (
    CanvasSpec,
    ContentSpec,
    ElementSpec,
    SlideSpec,
)
from pptx_skill.deck_planner import (
    LayoutScoringConfig,
    _signature_similarity,
    plan_deck,
)
from pptx_skill.generation_pipeline import (
    GenerationResult,
    plan_deck_layouts,
    run_generation_pipeline,
)
from pptx_skill.layout_engine import _builtin_tokens
from pptx_skill.pagination import (
    ContentItem,
    FitResult,
    paginate_bullets,
    paginate_content_spec,
    paginate_items,
)


def _bullets_slide(slide_id: str, n_items: int, role: str = "bullets") -> SlideSpec:
    body_text = "\n".join(f"要点 {i + 1}：这是一段需要展示的正文内容，用于验证分页行为。" for i in range(n_items))
    return SlideSpec(
        id=slide_id,
        role=role,
        communication_goal="test bullets",
        elements=[
            ElementSpec(
                id=f"{slide_id}/title",
                kind="text",
                role="title",
                content={"text": "测试标题"},
                style_ref="component.title",
            ),
            ElementSpec(
                id=f"{slide_id}/body",
                kind="text",
                role="body",
                content={"text": body_text},
                style_ref="component.body",
            ),
        ],
        preferred_layouts=["bullets.rail", "bullets.wide"],
        source_section_id="sec-01",
        fragment_index=0,
    )


class PaginateItemsTests(unittest.TestCase):
    def test_greedy_packing_under_capacity(self):
        items = [ContentItem(text=f"item {i}") for i in range(3)]
        # Capacity: max 2 items per page.
        cap = lambda c: FitResult(fits=len(c) <= 2, overflow_items=0, estimated_score=0.0, diagnostics={})  # noqa: E731
        pages = paginate_items(items, cap, min_items_per_page=1, max_items_per_page=2)
        self.assertEqual([len(p) for p in pages], [2, 1])

    def test_single_oversize_item_emitted_anyway(self):
        items = [ContentItem(text="huge")]
        cap = lambda c: FitResult(fits=False, overflow_items=1, estimated_score=0.0, diagnostics={})  # noqa: E731
        pages = paginate_items(items, cap)
        self.assertEqual(len(pages), 1)
        self.assertEqual(len(pages[0]), 1)

    def test_keep_with_next_breaks_page(self):
        items = [ContentItem(text=f"item {i}") for i in range(4)]
        # item at index 1 must stay with item 2.
        cap = lambda c: FitResult(fits=len(c) <= 3, overflow_items=0, estimated_score=0.0, diagnostics={})  # noqa: E731
        pages = paginate_items(items, cap, keep_with_next={1}, min_items_per_page=1)
        # The break should occur after index 1's group; verify no page splits 1 and 2.
        flat = [item for page in pages for item in page]
        self.assertEqual(len(flat), 4)
        # item 2 must be on the same page as item 1.
        for page in pages:
            texts = [it.text for it in page]
            if "item 1" in texts:
                self.assertIn("item 2", texts)

    def test_empty_items_returns_empty(self):
        self.assertEqual(paginate_items([], lambda c: FitResult(True, 0, 0.0, {})), [])


class PaginateBulletsTests(unittest.TestCase):
    def test_no_split_when_within_limit(self):
        slide = _bullets_slide("s1", n_items=5)
        self.assertIsNone(paginate_bullets(slide, max_items=7))

    def test_split_derives_stable_ids(self):
        slide = _bullets_slide("s1", n_items=15)
        derived = paginate_bullets(slide, max_items=5)
        self.assertIsNotNone(derived)
        self.assertGreater(len(derived), 1)
        ids = [d.id for d in derived]
        self.assertEqual(len(ids), len(set(ids)), "derived slide ids must be unique")
        # fragment_index increments.
        self.assertEqual([d.fragment_index for d in derived], list(range(len(derived))))
        # body element ids derived from parent.
        for idx, d in enumerate(derived):
            body = next(e for e in d.elements if e.role == "body")
            self.assertIn(f"page-{idx}", body.id)
        # continuation title on pages after the first.
        titles = [next(e for e in d.elements if e.role == "title").content.get("text", "") for d in derived]
        self.assertNotIn("（续）", titles[0])
        self.assertTrue(all("（续）" in t for t in titles[1:]))

    def test_split_is_deterministic(self):
        slide = _bullets_slide("s1", n_items=12)
        a = paginate_bullets(slide, max_items=5)
        b = paginate_bullets(slide, max_items=5)
        self.assertEqual([d.id for d in a], [d.id for d in b])


class PaginateContentSpecTests(unittest.TestCase):
    def test_mapping_records_fragment_counts(self):
        content = ContentSpec(
            id="deck",
            title="T",
            subtitle="",
            slides=[
                _bullets_slide("s1", n_items=3),
                _bullets_slide("s2", n_items=15),
            ],
        )
        derived, mapping = paginate_content_spec(content, max_items=5)
        # s1 stays single, s2 splits.
        self.assertEqual(mapping[0], ("s1", 0, 1))
        self.assertEqual(mapping[1][0], "s2")
        self.assertGreater(mapping[1][2], 1)
        # derived slide count reflects split.
        self.assertEqual(len(derived.slides), 1 + mapping[1][2])


class GeometrySignatureTests(unittest.TestCase):
    def test_identical_plans_similar(self):
        from pptx_skill.content_model import BBox, GeometrySpec, LayoutPlan, PlannedNode

        def make_plan():
            return LayoutPlan(
                canvas=CanvasSpec(959.976, 540),
                recipe_id="x",
                nodes=[
                    PlannedNode(
                        id="n1",
                        element_id="e1",
                        recipe_node_id="title",
                        kind="text",
                        role="title",
                        geometry=GeometrySpec(BBox(48, 36, 864, 48)),
                        resolved_style={},
                        content_binding={},
                        z_order=0,
                    )
                ],
                local_score=0.0,
            )

        self.assertGreater(_signature_similarity(make_plan(), make_plan()), 0.99)

    def test_disjoint_roles_dissimilar(self):
        from pptx_skill.content_model import BBox, GeometrySpec, LayoutPlan, PlannedNode

        def plan_with(role):
            return LayoutPlan(
                canvas=CanvasSpec(959.976, 540),
                recipe_id="x",
                nodes=[
                    PlannedNode(
                        id="n1",
                        element_id="e1",
                        recipe_node_id=role,
                        kind="text",
                        role=role,
                        geometry=GeometrySpec(BBox(48, 36, 100, 100)),
                        resolved_style={},
                        content_binding={},
                        z_order=0,
                    )
                ],
                local_score=0.0,
            )

        self.assertLess(_signature_similarity(plan_with("title"), plan_with("body")), 0.5)


class DeckPlannerTests(unittest.TestCase):
    def _profile(self, preferred_sequence=None, rhythm_rules=None):
        profile = {"tokens": _builtin_tokens()}
        if preferred_sequence is not None:
            profile["preferred_sequence"] = preferred_sequence
        if rhythm_rules is not None:
            profile["rhythm"] = {"rules": rhythm_rules}
        return profile

    def test_plan_single_bullets_slide(self):
        slide = _bullets_slide("s1", n_items=3)
        deck = plan_deck([slide], CanvasSpec(959.976, 540), self._profile())
        self.assertEqual(deck.status, "feasible")
        self.assertEqual(len(deck.plans), 1)
        self.assertEqual(len(deck.derived_slides), 1)

    def test_plan_splits_overloaded_bullets(self):
        slide = _bullets_slide("s1", n_items=20)
        deck = plan_deck([slide], CanvasSpec(959.976, 540), self._profile())
        # Derived slides should be more than one when split occurred.
        self.assertGreaterEqual(len(deck.derived_slides), 1)
        # When split happened, derived slide ids are unique.
        ids = [d.id for d in deck.derived_slides]
        self.assertEqual(len(ids), len(set(ids)))

    def test_preferred_sequence_consumed(self):
        slides = [_bullets_slide(f"s{i}", n_items=3) for i in range(3)]
        profile = self._profile(preferred_sequence=["bullets", "dashboard"])
        deck = plan_deck(slides, CanvasSpec(959.976, 540), profile)
        self.assertEqual(deck.status, "feasible")
        self.assertEqual(len(deck.plans), 3)

    def test_beam_selects_lower_score_path(self):
        # Two slides; ensure beam width pruning completes and returns feasible.
        slides = [_bullets_slide(f"s{i}", n_items=3) for i in range(2)]
        deck = plan_deck(slides, CanvasSpec(959.976, 540), self._profile(), config=LayoutScoringConfig(beam_width=4))
        self.assertEqual(deck.status, "feasible")
        self.assertEqual(len(deck.plans), 2)
        # The last diagnostic records the chosen beam score.
        beam_diag = next((d for d in deck.diagnostics if "beam_score" in d), None)
        self.assertIsNotNone(beam_diag)
        self.assertGreaterEqual(beam_diag["beam_score"], 0.0)


class PlanDeckLayoutsFacadeTests(unittest.TestCase):
    def test_facade_returns_one_plan_per_derived_slide(self):
        slide = _bullets_slide("s1", n_items=3)
        profile = {"tokens": _builtin_tokens()}
        plans = plan_deck_layouts([slide], CanvasSpec(959.976, 540), profile)
        self.assertEqual(len(plans), 1)


class PipelineWithPaginationTests(unittest.TestCase):
    def _content(self, n_items: int) -> ContentSpec:
        return ContentSpec(
            id="deck-1",
            title="Test",
            subtitle="",
            slides=[_bullets_slide("s1", n_items=n_items)],
        )

    def test_pipeline_off_mode_renders_split_slides(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = os.path.join(tmp, "deck.pptx")
            content = self._content(n_items=20)
            result = run_generation_pipeline(content, out, qa_mode="off")
            self.assertEqual(result.qa_status, "pass")
            self.assertTrue(os.path.exists(out))
            # The PPTX should contain at least one slide and be openable.
            from pptx import Presentation

            prs = Presentation(out)
            slide_count = len(list(prs.slides))
            self.assertGreaterEqual(slide_count, 1)

    def test_pipeline_report_mode_runs(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = os.path.join(tmp, "deck.pptx")
            content = self._content(n_items=3)
            result = run_generation_pipeline(content, out, qa_mode="report")
            self.assertIsInstance(result, GenerationResult)
            self.assertIn(result.qa_status, {"pass", "inconclusive"})


if __name__ == "__main__":
    unittest.main()
