"""Tests for the api.py compatibility facade."""
from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from pptx import Presentation

from pptx_skill import (
    GenerationResult,
    auto_generate_ppt,
    auto_validate_ppt,
)
from pptx_skill.visual_qa import CheckOutcome


class ApiFacadeTests(unittest.TestCase):
    def test_legacy_string_return(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "deck.pptx"
            result = auto_generate_ppt(
                title="兼容",
                sections=[{"title": "页", "bullets": ["1", "2"]}],
                output_path=str(path),
                theme_key="editorial",
                auto_search_images=False,
            )
            self.assertIsInstance(result, str)
            self.assertTrue(Path(result).exists())

    def test_return_result_structured(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "deck.pptx"
            result = auto_generate_ppt(
                title="结构化",
                sections=[{"title": "页", "bullets": ["1"]}],
                output_path=str(path),
                theme_key="editorial",
                auto_search_images=False,
                qa_mode="report",
                return_result=True,
            )
            self.assertIsInstance(result, GenerationResult)
            self.assertEqual(result.qa_status, CheckOutcome.PASS)
            self.assertTrue(Path(result.pptx_path).exists())

    def test_auto_validate_legacy_format(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "deck.pptx"
            auto_generate_ppt(
                title="验证",
                sections=[{"title": "页", "bullets": ["1"]}],
                output_path=str(path),
                theme_key="editorial",
                auto_search_images=False,
            )
            report = auto_validate_ppt(str(path))
            self.assertIn("passed", report)
            self.assertIn("checks", report)
            self.assertIn("warnings", report)
            self.assertIn("total_slides", report)

    def test_auto_validate_new_report(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "deck.pptx"
            auto_generate_ppt(
                title="验证",
                sections=[{"title": "页", "bullets": ["1"]}],
                output_path=str(path),
                theme_key="editorial",
                auto_search_images=False,
            )
            report = auto_validate_ppt(str(path), return_report=True)
            self.assertEqual(report.status, CheckOutcome.PASS)

    def test_facade_validate_matches_legacy_checks(self):
        # The facade validate must run the same five checks as the legacy
        # implementation (no silent loss of font-hierarchy / color-restraint).
        import sys
        sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
        from pptx_helper import auto_validate_ppt as legacy_validate

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "deck.pptx"
            auto_generate_ppt(
                title="一致性",
                sections=[
                    {"title": "页一", "bullets": ["1", "2", "3"]},
                    {"title": "页二", "steps": ["发现", "设计", "交付"]},
                ],
                output_path=str(path),
                theme_key="editorial",
                auto_search_images=False,
            )
            legacy = legacy_validate(str(path))
            facade = auto_validate_ppt(str(path))
            self.assertEqual(
                [c[0] for c in facade["checks"]],
                [c[0] for c in legacy["checks"]],
            )
            self.assertEqual(facade["passed"], legacy["passed"])
            self.assertEqual(facade["total_slides"], legacy["total_slides"])


class AutoPageDedupTests(unittest.TestCase):
    """Caller-provided cover/toc/end sections must not be duplicated by the
    auto cover/toc/end pages in the legacy generator."""

    def _generate(self, tmp, sections):
        path = Path(tmp) / "deck.pptx"
        auto_generate_ppt(
            title="去重",
            sections=sections,
            output_path=str(path),
            theme_key="editorial",
            auto_search_images=False,
        )
        return len(Presentation(str(path)).slides)

    def test_explicit_end_not_duplicated(self):
        with tempfile.TemporaryDirectory() as tmp:
            n = self._generate(tmp, [
                {"title": "要点", "bullets": ["1", "2"], "layout": "bullets"},
                {"title": "谢谢", "layout": "end"},
            ])
            self.assertEqual(n, 3)  # cover + bullets + end

    def test_auto_end_still_added(self):
        with tempfile.TemporaryDirectory() as tmp:
            n = self._generate(tmp, [
                {"title": "要点", "bullets": ["1"], "layout": "bullets"},
            ])
            self.assertEqual(n, 3)  # cover + bullets + end

    def test_explicit_cover_not_duplicated(self):
        with tempfile.TemporaryDirectory() as tmp:
            n = self._generate(tmp, [
                {"title": "自定义封面", "subtitle": "副", "layout": "cover"},
                {"title": "要点", "bullets": ["1"], "layout": "bullets"},
            ])
            self.assertEqual(n, 3)  # cover + bullets + end

    def test_explicit_toc_not_duplicated(self):
        with tempfile.TemporaryDirectory() as tmp:
            n = self._generate(tmp, [
                {"title": "目录", "layout": "toc", "bullets": ["一", "二", "三", "四"]},
                *[{"title": f"S{i}", "bullets": ["1"], "layout": "bullets"} for i in range(4)],
            ])
            self.assertEqual(n, 7)  # cover + toc + 4 + end

    def test_auto_toc_still_added(self):
        with tempfile.TemporaryDirectory() as tmp:
            n = self._generate(tmp, [
                *[{"title": f"S{i}", "bullets": ["1"], "layout": "bullets"} for i in range(4)],
            ])
            self.assertEqual(n, 7)  # cover + toc + 4 + end

    def test_poster_toc_keeps_nine_section_titles(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "deck.pptx"
            titles = [f"章节{i}" for i in range(1, 10)]
            auto_generate_ppt(
                title="长目录",
                sections=[
                    {"title": title, "bullets": ["1"], "layout": "bullets"}
                    for title in titles
                ],
                output_path=str(path),
                template_key="creative-editorial",
                auto_search_images=False,
            )
            presentation = Presentation(str(path))
            toc_text = "\n".join(
                shape.text
                for shape in presentation.slides[1].shapes
                if getattr(shape, "has_text_frame", False)
            )
            for title in titles:
                self.assertIn(title, toc_text)

    def test_auto_toc_excludes_explicit_cover(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "deck.pptx"
            auto_generate_ppt(
                title="封面标题",
                sections=[
                    {"title": "封面标题", "layout": "cover"},
                    *[
                        {"title": f"内容{i}", "bullets": ["1"], "layout": "bullets"}
                        for i in range(1, 5)
                    ],
                ],
                output_path=str(path),
                template_key="creative-editorial",
                auto_search_images=False,
            )
            presentation = Presentation(str(path))
            cover_text = "\n".join(
                shape.text
                for shape in presentation.slides[0].shapes
                if getattr(shape, "has_text_frame", False)
            )
            self.assertIn("封面标题", cover_text)
            toc_text = "\n".join(
                shape.text
                for shape in presentation.slides[1].shapes
                if getattr(shape, "has_text_frame", False)
            )
            self.assertNotIn("封面标题", toc_text)
            for index in range(1, 5):
                self.assertIn(f"内容{index}", toc_text)

    def test_poster_timeline_has_room_for_daypart_label(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "deck.pptx"
            auto_generate_ppt(
                title="路线",
                sections=[
                    {
                        "title": "48 小时路线",
                        "layout": "timeline",
                        "events": [{"date": "DAY 1 · AM", "title": "出发"}],
                    }
                ],
                output_path=str(path),
                template_key="creative-editorial",
                auto_search_images=False,
            )
            presentation = Presentation(str(path))
            labels = [
                shape
                for shape in presentation.slides[1].shapes
                if getattr(shape, "has_text_frame", False) and shape.text == "DAY 1 · AM"
            ]
            self.assertEqual(len(labels), 1)
            self.assertGreaterEqual(labels[0].width, 2.5 * 914400)


if __name__ == "__main__":
    unittest.main()
