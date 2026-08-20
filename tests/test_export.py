"""Regression tests for export range handling and temporary-resource cleanup."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from pptx import Presentation
from pptx.util import Inches

from pptx_skill.export import (
    export_thumbnails,
    export_to_images,
    export_to_pdf,
)


def _make_deck(path: Path, titles: list[str]) -> Path:
    prs = Presentation()
    for title in titles:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        textbox.text = title
    prs.save(path)
    return path


class ExportRegressionTests(unittest.TestCase):
    def test_pdf_range_drops_relationship_ids_not_numeric_slide_ids(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = _make_deck(root / "source.pptx", ["First", "Middle", "Last"])
            temp_dirs: list[Path] = []

            def inspect_trimmed_deck(pptx_path: str, output_path: str, dpi: int) -> str:
                trimmed = Presentation(pptx_path)
                self.assertEqual(len(trimmed.slides), 1)
                self.assertEqual(trimmed.slides[0].shapes[0].text, "Middle")
                temp_dirs.append(Path(pptx_path).parent)
                Path(output_path).write_bytes(b"%PDF-1.4\n%%EOF\n")
                return output_path

            with patch("pptx_skill.export._convert_to_pdf", side_effect=inspect_trimmed_deck):
                result = export_to_pdf(source, root / "middle.pdf", range=(2, 2))
            self.assertTrue(Path(result).exists())
            self.assertEqual(len(temp_dirs), 1)
            self.assertFalse(temp_dirs[0].exists())

    def test_render_temp_directories_are_removed_when_rendering_fails(self):
        exporters = (
            ("images", export_to_images),
            ("thumbnails", export_thumbnails),
        )
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = _make_deck(root / "source.pptx", ["Slide"])
            for name, exporter in exporters:
                with self.subTest(exporter=name):
                    render_dir = root / f"{name}-render"

                    def make_render_dir(
                        *args: object,
                        target: Path = render_dir,
                        **kwargs: object,
                    ) -> str:
                        target.mkdir()
                        return str(target)

                    with (
                        patch("pptx_skill.export.tempfile.mkdtemp", side_effect=make_render_dir),
                        patch("pptx_skill.export.render_preview", side_effect=RuntimeError("renderer failed")),
                    ):
                        with self.assertRaisesRegex(RuntimeError, "renderer failed"):
                            exporter(source, root / f"{name}-output")
                    self.assertFalse(render_dir.exists())


if __name__ == "__main__":
    unittest.main()
