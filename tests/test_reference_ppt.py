import sys
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS))

from reference_ppt import (  # noqa: E402
    analyze_presentation,
    compose_from_reference,
    generate_from_reference,
)
from template_engine import generate_template_preview, load_template_profile  # noqa: E402


class ReferencePptTests(unittest.TestCase):
    def _make_reference(self, directory: Path) -> Path:
        path = directory / "reference.pptx"
        generate_template_preview(load_template_profile("strategy-consulting"), path)
        return path

    def test_analyze_and_exact_shape_replacement(self):
        with tempfile.TemporaryDirectory() as tmp:
            directory = Path(tmp)
            reference = self._make_reference(directory)
            analysis = analyze_presentation(reference)
            self.assertEqual(analysis["slide_count"], 9)
            self.assertEqual(analysis["recommended_mode"], "clone")

            reference_deck = Presentation(str(reference))
            cover_text_shapes = [
                shape for shape in reference_deck.slides[0].shapes
                if getattr(shape, "has_text_frame", False) and shape.text.strip()
            ]
            end_text_shapes = [
                shape for shape in reference_deck.slides[-1].shapes
                if getattr(shape, "has_text_frame", False) and shape.text.strip()
            ]
            self.assertGreaterEqual(len(cover_text_shapes), 2)
            self.assertGreaterEqual(len(end_text_shapes), 1)

            output = directory / "composed.pptx"
            compose_from_reference(reference, {
                "slides": [
                    {"source_slide": 1, "replacements": {
                        cover_text_shapes[0].name: {"text": "全新标题"},
                        cover_text_shapes[1].name: {"text": "严格沿用参考设计"}
                    }},
                    {"source_slide": 9, "replacements": {
                        end_text_shapes[0].name: {"text": "完成"}
                    }}
                ]
            }, output)
            result = Presentation(str(output))
            self.assertEqual(len(result.slides), 2)
            all_text = "\n".join(
                shape.text for slide in result.slides for shape in slide.shapes
                if getattr(shape, "has_text_frame", False)
            )
            self.assertIn("全新标题", all_text)
            self.assertIn("严格沿用参考设计", all_text)
            self.assertIn("完成", all_text)

    def test_native_and_clone_generation_modes(self):
        with tempfile.TemporaryDirectory() as tmp:
            directory = Path(tmp)
            reference = self._make_reference(directory)
            sections = [{"title": "执行路径", "steps": ["分析", "设计", "上线"]}]
            for mode in ("native", "clone"):
                output = directory / f"{mode}.pptx"
                generate_from_reference(
                    reference, "客户项目", sections, output,
                    subtitle="参考模板复刻", mode=mode,
                )
                result = Presentation(str(output))
                self.assertEqual(len(result.slides), 3)
                text = "\n".join(
                    shape.text for slide in result.slides for shape in slide.shapes
                    if getattr(shape, "has_text_frame", False)
                )
                self.assertIn("客户项目", text)


if __name__ == "__main__":
    unittest.main()
