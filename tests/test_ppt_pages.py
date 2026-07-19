import os
import tempfile
import unittest

from pptx import Presentation

from pptx_skill import auto_generate_ppt, delete_slide, insert_slide


class TestDeleteInsertRoundTrip(unittest.TestCase):
    def test_delete_then_insert_no_duplicate_partname(self):

        with tempfile.TemporaryDirectory() as tmp:
            pptx_path = os.path.join(tmp, "test.pptx")
            sections = [
                {"title": "第一章", "bullets": ["要点一", "要点二"]},
                {"title": "第二章", "bullets": ["要点三", "要点四"]},
                {"title": "第三章", "bullets": ["要点五", "要点六"]},
                {"title": "第四章", "bullets": ["要点七", "要点八"]},
            ]
            auto_generate_ppt(
                title="测试",
                sections=sections,
                output_path=pptx_path,
                auto_search_images=False,
            )
            n_after_gen = len(Presentation(pptx_path).slides)
            self.assertGreaterEqual(n_after_gen, 3)

            delete_slide(pptx_path, 2)
            n_after_del = len(Presentation(pptx_path).slides)
            self.assertEqual(n_after_del, n_after_gen - 1)

            insert_slide(
                pptx_path, 2,
                {"title": "新插入页", "bullets": ["新要点"]},
                auto_search_images=False,
            )
            n_after_ins = len(Presentation(pptx_path).slides)
            self.assertEqual(n_after_ins, n_after_gen)

            prs = Presentation(pptx_path)
            partnames = [str(p.partname) for p in prs.part.package.iter_parts()]
            self.assertEqual(len(partnames), len(set(partnames)),
                             f"Duplicate partnames found: {[p for p in partnames if partnames.count(p) > 1]}")

            save_path = os.path.join(tmp, "resaved.pptx")
            prs.save(save_path)
            prs2 = Presentation(save_path)
            self.assertEqual(len(prs2.slides), n_after_ins)
