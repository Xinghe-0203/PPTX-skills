"""Regression tests for shared presentation I/O helpers."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from pptx import Presentation

from pptx_skill._io import ensure_path_on_disk, save_prs


class _PartialWriter:
    def save(self, path: str | Path) -> None:
        Path(path).write_bytes(b"PARTIAL")
        raise RuntimeError("simulated save failure")


class _BytesWriter:
    def __init__(self, data: bytes) -> None:
        self.data = data

    def save(self, path: str | Path) -> None:
        Path(path).write_bytes(self.data)


class SharedIoTests(unittest.TestCase):
    def test_save_failure_preserves_original_and_cleans_temp_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            target = root / "deck.pptx"
            target.write_bytes(b"ORIGINAL")

            with self.assertRaisesRegex(RuntimeError, "simulated save failure"):
                save_prs(_PartialWriter(), target)

            self.assertEqual(target.read_bytes(), b"ORIGINAL")
            self.assertEqual(target.with_suffix(".bak.pptx").read_bytes(), b"ORIGINAL")
            self.assertEqual(
                {item.name for item in root.iterdir()},
                {"deck.pptx", "deck.bak.pptx"},
            )

    def test_successful_save_replaces_target_and_keeps_backup(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            target = root / "deck.pptx"
            original = Presentation()
            original.slides.add_slide(original.slide_layouts[6])
            original.save(target)

            replacement = Presentation()
            replacement.slides.add_slide(replacement.slide_layouts[6])
            replacement.slides.add_slide(replacement.slide_layouts[6])
            save_prs(replacement, target)

            self.assertEqual(len(Presentation(target).slides), 2)
            self.assertEqual(len(Presentation(target.with_suffix(".bak.pptx")).slides), 1)

    def test_replace_failure_preserves_original_and_cleans_temp_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            target = root / "deck.pptx"
            target.write_bytes(b"ORIGINAL")

            with (
                patch("pptx_skill._io.os.replace", side_effect=OSError("replace failed")),
                self.assertRaisesRegex(OSError, "replace failed"),
            ):
                save_prs(_BytesWriter(b"REPLACEMENT"), target, backup=False)

            self.assertEqual(target.read_bytes(), b"ORIGINAL")
            self.assertEqual([item.name for item in root.iterdir()], ["deck.pptx"])

    def test_ensure_path_on_disk_cleans_temp_directory_on_save_failure(self):
        with tempfile.TemporaryDirectory() as tmp:
            temp_dir = Path(tmp) / "materialized"

            def make_temp_dir(*args: object, **kwargs: object) -> str:
                temp_dir.mkdir()
                return str(temp_dir)

            with (
                patch("pptx_skill._io.resolve_path", return_value=None),
                patch("pptx_skill._io.open_prs", return_value=_PartialWriter()),
                patch("pptx_skill._io.tempfile.mkdtemp", side_effect=make_temp_dir),
                self.assertRaisesRegex(RuntimeError, "simulated save failure"),
            ):
                ensure_path_on_disk(object())

            self.assertFalse(temp_dir.exists())


if __name__ == "__main__":
    unittest.main()
