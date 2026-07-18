"""Tests for the adaptive PPTX renderer."""
from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from pptx_skill.content_model import (
    BBox,
    CanvasSpec,
    GeometrySpec,
    LayoutPlan,
    PlannedNode,
    SafeInsets,
)
from pptx_skill.pptx_renderer import render_layout_plan


class PptxRendererTests(unittest.TestCase):
    def _make_image(self, directory: Path) -> Path:
        from PIL import Image

        path = directory / "img.png"
        Image.new("RGB", (400, 300), color=(0, 128, 0)).save(path)
        return path

    def test_render_text_and_image_nodes(self):
        with tempfile.TemporaryDirectory() as tmp:
            img_path = self._make_image(Path(tmp))
            out = Path(tmp) / "adaptive.pptx"
            canvas = CanvasSpec(959.976, 540, safe=SafeInsets(36, 48, 32, 48))
            plan = LayoutPlan(
                canvas=canvas,
                recipe_id="test.bullets",
                nodes=[
                    PlannedNode(
                        id="s0/title",
                        element_id="title",
                        recipe_node_id="title",
                        kind="text",
                        role="title",
                        geometry=GeometrySpec(BBox(100, 50, 800, 80)),
                        resolved_style={"size": 32, "color": "#1A1A1A"},
                        content_binding={"text": "自适应标题"},
                        z_order=1,
                    ),
                    PlannedNode(
                        id="s0/hero",
                        element_id="hero",
                        recipe_node_id="hero",
                        kind="image",
                        role="hero",
                        geometry=GeometrySpec(BBox(560, 120, 340, 260)),
                        resolved_style={},
                        content_binding={"path": str(img_path)},
                        z_order=0,
                    ),
                ],
                local_score=0.0,
            )
            result = render_layout_plan(plan, str(out))
            self.assertTrue(Path(result.pptx_path).exists())
            self.assertEqual(len(result.trace), 2)

            prs = Presentation(str(out))
            self.assertEqual(len(prs.slides), 1)
            self.assertAlmostEqual(
                float(prs.slide_width / Inches(1)), canvas.width_pt / 72, places=2
            )

    def test_render_shape_node(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "shape.pptx"
            canvas = CanvasSpec(959.976, 540)
            plan = LayoutPlan(
                canvas=canvas,
                recipe_id="test.shape",
                nodes=[
                    PlannedNode(
                        id="s0/bg",
                        element_id=None,
                        recipe_node_id="bg",
                        kind="shape",
                        role="decoration",
                        geometry=GeometrySpec(BBox(0, 0, 200, 200)),
                        resolved_style={"fill": "#D9623B"},
                        content_binding={"shape_type": "rectangle"},
                        z_order=0,
                        decorative=True,
                    )
                ],
                local_score=0.0,
            )
            result = render_layout_plan(plan, str(out))
            self.assertEqual(len(result.trace), 1)
            prs = Presentation(str(out))
            self.assertEqual(len(prs.slides[0].shapes), 1)


if __name__ == "__main__":
    unittest.main()
